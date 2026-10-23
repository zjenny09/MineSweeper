#!/usr/bin/env python3
"""Maintain stable art IDs, metadata, the review workbook, and Godot lookups.

Commands:
  initialize  Create the first ART_#### catalog from authoritative assets.
  discover    Preserve known IDs and append IDs for newly discovered assets.
  export      Refresh metadata and rebuild Excel/Godot generated files.
  import      Apply editable Excel fields to JSON and rebuild Godot lookups.
  validate    Check IDs, names, paths, metadata, workbook, and source usage.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "docs" / "art" / "GAME_ART_CATALOG.json"
WORKBOOK = ROOT / "docs" / "art" / "GAME_ART_CATALOG.xlsx"
GENERATED_GDSCRIPT = ROOT / "scripts" / "art_ids.gd"
ART_CATALOG = ROOT / "scripts" / "art_catalog.gd"
PROMPT_DOCUMENT = ROOT / "docs" / "art" / "AI_ART_PROMPT_SYSTEM.md"
ID_PATTERN = re.compile(r"^ART_\d{4}$")
SHORT_NAME_PATTERN = re.compile(r"^[a-z][a-z0-9]*(?:_[a-z0-9]+)*$")
MAX_SHORT_NAME_LENGTH = 36
VISUAL_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".svg", ".ttf", ".otf"}
SHADER_EXTENSIONS = {".gdshader", ".shader"}
SOURCE_EXTENSIONS = {".gd", ".tscn", ".tres", ".godot", ".cfg"}
RUNTIME_ROOTS = ["scripts", "scenes", "themes", "animations"]

PROMPTS: list[tuple[str, str, str, str]] = [
    ("PRM_STYLE", "全局母提示词", "第5节", "所有素材的统一风格前缀"),
    ("PRM_STYLE_BOARD", "风格测试板", "第6节", "风格与缩略可读性确认"),
    ("PRM_ROBOT", "小芽机器人", "第7节", "主角、机器人和章节变体"),
    ("PRM_MONSTER", "污染核心史莱姆", "第8节", "钟罩形污染核心角色、状态与章节变体"),
    ("PRM_MENU_BG", "欢迎页背景", "第9.1节", "主菜单和欢迎页背景"),
    ("PRM_BUTTON", "按钮底板", "第9.2、9.3、22节", "按钮及交互状态"),
    ("PRM_PANEL", "纸卡与弹窗", "第10节", "信息卡、对话框和弹窗"),
    ("PRM_BG_LAND", "陆地背景", "第11.1节", "陆地章节背景"),
    ("PRM_BG_OCEAN", "海洋背景", "第11.2节", "海洋章节背景"),
    ("PRM_BG_SKY", "天空背景", "第11.3节", "天空章节背景"),
    ("PRM_BG_SPACE", "宇宙背景", "第11.4节", "宇宙纹理与星云 Shader"),
    ("PRM_STAGE_LAND", "陆地舞台", "第12.1节", "陆地棋盘框和舞台"),
    ("PRM_STAGE_OCEAN", "海洋舞台", "第12.2节", "海洋棋盘框和舞台"),
    ("PRM_STAGE_SKY", "天空舞台", "第12.3节", "天空棋盘框和舞台"),
    ("PRM_HUD_SPACE", "宇宙 HUD", "第12.4节", "宇宙信息卡"),
    ("PRM_CELL_LAND", "陆地方格", "第13.2节", "陆地棋盘格状态"),
    ("PRM_CELL_OCEAN", "海洋六边格", "第13.3节", "海洋棋盘格状态"),
    ("PRM_CELL_SKY", "天空云膜格", "第13.4节", "天空五边形和六边形格"),
    ("PRM_NODE_SPACE", "宇宙星核", "第13.5、18节", "宇宙节点和图语法"),
    ("PRM_MARKER", "标记与图标", "第14节", "标记、核心和操作图标"),
    ("PRM_MECH_LAND", "陆地机制", "第15节", "障碍、污染节点和古树"),
    ("PRM_MECH_OCEAN", "海洋机制", "第16节", "珊瑚墙、水雷和洋流"),
    ("PRM_MECH_SKY", "天空机制", "第17节", "微风、云盾和雷眼"),
    ("PRM_SPACE", "宇宙图语法", "第18节", "星核外观和宇宙飞船"),
    ("PRM_ECOLOGY", "生态主体", "第19、23节", "健康与失败生态状态"),
    ("PRM_MAP", "章节地图", "第20节", "地图底图和入口卡"),
    ("PRM_MAP_NODE", "地图节点", "第21节", "地图节点和状态"),
    ("PRM_EFFECT", "透明特效", "第24节", "光效、粒子与透明修正"),
    ("PRM_ALPHA", "透明素材修正", "第24节", "透明边缘清理"),
    ("PRM_FONT", "字体与文字规范", "第3.2、27节", "字体资产及文字不入图规则"),
]
PROMPT_IDS = {item[0] for item in PROMPTS}

WORKBOOK_HEADERS = [
    "代号",
    "规范短名",
    "素材名称（中文）",
    "内容说明",
    "章节/范围",
    "类别",
    "状态/变体",
    "当前路径",
    "文件格式",
    "尺寸",
    "透明通道",
    "文件大小KB",
    "SHA-256",
    "运行状态",
    "程序别名",
    "使用位置（仅供参考）",
    "提示词代号",
    "重制状态",
    "权利状态",
    "来源/生成工具",
    "商用凭证",
    "修改备注",
    "审核状态",
]
EDITABLE_COLUMNS = {
    "规范短名": "short_name",
    "素材名称（中文）": "asset_name",
    "内容说明": "description",
    "章节/范围": "scope",
    "类别": "category",
    "状态/变体": "variant",
    "当前路径": "path",
    "提示词代号": "prompt_id",
    "重制状态": "replacement_status",
    "权利状态": "rights_status",
    "来源/生成工具": "source_tool",
    "商用凭证": "commercial_evidence",
    "修改备注": "note",
    "审核状态": "review",
}

SCOPE_PREFIXES = {
    "common": ("通用", "cmn"),
    "gameplay": ("玩法通用", "play"),
    "tutorial": ("教学", "tut"),
    "welcome": ("欢迎页", "wel"),
    "level_01": ("陆地", "land"),
    "land_levels": ("陆地", "land"),
    "level_select": ("章节地图", "map"),
    "ocean_levels": ("海洋", "ocean"),
    "sky_levels": ("天空", "sky"),
}

VARIANT_NAMES = {
    "normal": "常态",
    "hover": "悬停",
    "pressed": "按下",
    "focus": "键盘焦点",
    "hidden": "未翻开",
    "revealed": "已翻开",
    "polluted": "污染",
    "healthy": "健康",
    "wilted": "衰败",
    "failed": "失败",
    "wrong": "误标",
    "active": "激活",
    "cleansed": "已净化",
    "clean": "清洁版",
    "crying": "哭泣",
    "smiling": "微笑",
    "standing": "站立",
    "sitting": "坐姿",
    "upright": "直立",
    "front": "正面",
    "back": "背面",
    "right": "右侧",
    "left": "左侧",
    "planting": "种植",
    "chasing": "追逐",
    "squashed": "压扁",
    "rebuilt": "重制版",
    "grayblue": "灰蓝版",
    "dark": "深色版",
    "uprooted": "拔起错误态",
}

ASSET_STATE_SUFFIXES = [
    ("keyboard_focus", "键盘焦点"),
    ("facing_right", "面向右"),
    ("facing_left", "面向左"),
    ("light_green", "浅绿色"),
    ("dark_green", "深绿色"),
    ("healthy", "健康"),
    ("polluted", "污染"),
    ("revealed", "已翻开"),
    ("cleansed", "已净化"),
    ("pressed", "按下"),
    ("standing", "站立"),
    ("squashed", "压扁"),
    ("smiling", "微笑"),
    ("planting", "种植"),
    ("chasing", "追逐"),
    ("upright", "直立"),
    ("wilted", "衰败"),
    ("failed", "失败"),
    ("normal", "常态"),
    ("hover", "悬停"),
    ("focus", "键盘焦点"),
    ("hidden", "未翻开"),
    ("active", "激活"),
    ("crying", "哭泣"),
    ("sitting", "坐姿"),
    ("front", "正面"),
    ("back", "背面"),
    ("right", "右侧"),
    ("left", "左侧"),
    ("clean", "清洁版"),
    ("wrong", "错误"),
    ("rebuilt", "重制版"),
    ("grayblue", "灰蓝版"),
    ("dark", "深色版"),
    ("uprooted", "拔起错误态"),
]

ASSET_SUBJECT_NAMES = {
    "ui_handwritten_zh": "中文手写界面字体",
    "patrickhand-regular": "英文 Patrick Hand 手写字体",
    "zcoolkuaile-regular": "中文站酷快乐体字体",
    "universe_nebula": "宇宙星云背景着色器",
    "menu_primary": "主菜单主按钮",
    "menu_secondary": "主菜单次按钮",
    "first_move_button_primary": "首步引导主按钮",
    "first_move_button_secondary": "首步引导次按钮",
    "first_move_guide_frame": "首步引导对话框底板",
    "first_move_guide_robot_avatar": "首步引导机器人头像",
    "flag_sprout": "陆地小芽标记",
    "pollution_core_slime": "通用污染核心史莱姆",
    "land_paper_background": "第1关纸张背景",
    "land_tabletop_background": "陆地桌面背景（旧版）",
    "land_tabletop_background5": "陆地桌面背景（当前版）",
    "pause": "陆地暂停按钮",
    "regenerate": "陆地重新生成按钮",
    "round_action": "第1关圆形操作按钮",
    "cell": "陆地方形棋盘格",
    "board_tray": "第1关棋盘托盘",
    "board_shadow": "第1关棋盘阴影",
    "unified_land_stage_body": "陆地统一棋盘舞台主体",
    "unified_land_stage_shadow": "陆地统一棋盘舞台阴影",
    "note_curved": "第1关弧形信息便签",
    "note_straight": "第1关直边信息便签",
    "pause_menu_frame": "第1关暂停菜单框",
    "background_seed_sticker": "第1关背景种子贴纸",
    "background_leaf_sprout_sticker": "第1关背景叶芽贴纸",
    "background_green_leaf": "第1关背景绿叶",
    "left_panel_sprout": "第1关左侧小芽生态",
    "status_note_bud_sticker": "第1关状态便签花苞贴纸",
    "left_guardian": "陆地左侧守护者",
    "left_robot_guardian": "陆地左侧机器人守护者",
    "left_sprout_guardian": "第1关左侧小芽守护者",
    "right_panel_guardian_white": "第1关右侧白色守护者",
    "right_panel_guardian_yellow": "第1关右侧黄色守护者",
    "right_white_guardian": "陆地右侧白色守护者",
    "right_yellow_guardian": "陆地右侧黄色守护者",
    "top_guardian_white": "第1关顶部白色守护者",
    "top_guardian_yellow": "第1关顶部黄色守护者",
    "top_guardian_light_green": "第1关顶部浅绿守护者",
    "top_guardian_dark_green": "第1关顶部深绿守护者",
    "robot_failure_body": "陆地机器人失败动作主体",
    "robot_reaction_body_with_sprout": "陆地机器人带小芽反应主体",
    "tabletop_paper_monster": "陆地桌面污染纸怪",
    "level_marker_button": "陆地地图关卡节点按钮",
    "ocean_level_marker_star": "海洋地图普通星形节点",
    "ocean_level_marker_final": "海洋地图终点节点",
    "shortcut_button": "章节地图快捷入口按钮",
    "bottom_info_bar": "章节地图底部信息栏",
    "desktop_background": "陆地选关桌面背景（旧版）",
    "desktop_background2": "陆地选关桌面背景（当前版）",
    "land_map": "陆地章节地图（旧版）",
    "land_map2": "陆地章节地图（当前版）",
    "land_map_tilted_shadow": "陆地章节地图倾斜阴影",
    "ocean_level_map": "海洋章节地图",
    "ocean_level_select_background": "海洋选关背景",
    "sky_level_map": "天空章节地图",
    "ocean_pause": "海洋暂停按钮",
    "ocean_regenerate": "海洋重新生成按钮",
    "ocean_desktop": "海洋游戏桌面背景",
    "ocean_board_tray_frame": "海洋棋盘托盘框",
    "ocean_cell": "海洋六边形棋盘格",
    "ocean_reef_strip": "海洋连续珊瑚礁墙分段",
    "ocean_reef_strip_full": "海洋完整珊瑚礁墙",
    "ocean_flag_coral": "海洋珊瑚标记",
    "ocean_marker_monster": "海洋错误污染怪标记",
    "ocean_pollution_core_monster": "海洋污染核心怪",
    "ocean_robot_helmet": "海洋主角机器人头盔",
    "ocean_robot_reaction_body": "海洋主角机器人反应主体",
    "ocean_bubble_monster": "海洋气泡污染怪",
    "sky_desktop": "天空游戏桌面背景",
    "sky_board_frame": "天空球面棋盘框",
    "sky_flag_plane": "天空纸飞机标记",
    "sky_marker_wrong_cloud": "天空错误标记云",
    "sky_pollution_core_cloud": "天空污染核心云",
    "hexagon_surface": "天空六边形棋盘格",
    "hexagon_hidden_surface": "天空六边形棋盘格（未翻开）",
    "hexagon_revealed_surface": "天空六边形棋盘格（已翻开）",
    "hexagon_polluted_surface": "天空六边形棋盘格（污染）",
    "pentagon_surface": "天空五边形棋盘格",
    "pentagon_hidden_surface": "天空五边形棋盘格（未翻开）",
    "pentagon_revealed_surface": "天空五边形棋盘格（已翻开）",
    "pentagon_polluted_surface": "天空五边形棋盘格（污染）",
    "sky_robot": "天空主角机器人",
    "sky_robot_plane": "天空主角机器人飞行器",
    "sky_robot_paper_plane": "天空主角机器人纸飞机",
    "sky_robot_reaction_body": "天空主角机器人反应主体",
    "sky_robot_victory_body": "天空主角机器人胜利主体",
    "sky_robot_failure_body": "天空主角机器人失败主体",
    "sky_monster_paper_plane": "天空污染怪纸飞机",
    "landscape_base": "欢迎页山谷背景底图",
    "sun_glow": "欢迎页太阳光晕",
    "pollen_particles": "欢迎页花粉粒子",
    "water_glint": "欢迎页水面闪光",
    "robot": "欢迎页机器人",
    "robot_front": "欢迎页机器人正面",
    "robot_head_sprout": "欢迎页机器人头顶小芽",
    "robot_arm": "欢迎页机器人手臂",
    "robot_shoulder": "欢迎页机器人肩部",
    "slime": "欢迎页污染史莱姆",
    "slime_chase_shell": "欢迎页追逐史莱姆外壳",
    "slime_chase_contents": "欢迎页追逐史莱姆内部",
}

LAND_LEVEL_SUBJECTS = {
    "shrub": "灌木生态",
    "wetland": "湿地生态",
    "grass": "草地生态",
    "tree": "树木生态",
    "forest_edge": "森林边缘生态",
    "ancient_forest": "古老森林生态",
    "boss_tree": "Boss古树",
    "sticker": "装饰贴纸",
    "obstacle_broken_wall": "断墙障碍",
    "obstacle_web": "蛛网障碍",
    "obstacle_mossy_stone": "苔藓石障碍",
    "short_wall": "短墙障碍",
    "pollution_node": "污染节点",
}

OCEAN_LEVEL_SUBJECTS = {
    "anemone": "海葵生态",
    "seagrass": "海草床生态",
    "coral": "珊瑚生态",
    "kelp": "海带生态",
    "whale_fall": "鲸落生态",
    "jellyfish": "水母生态",
    "fish_school": "鱼群生态",
    "trench": "海沟生态",
}

SKY_LEVEL_SUBJECTS = {
    15: "云冠生态",
    16: "风环生态",
    17: "虹桥生态",
    18: "浮岛生态",
    19: "云脊生态",
    20: "雷眼生态",
    21: "天宫生态",
}


def res_path(path: Path) -> str:
    return "res://" + path.relative_to(ROOT).as_posix()


def local_path(value: str) -> Path:
    if not value.startswith("res://"):
        raise ValueError(f"Asset path must start with res://: {value}")
    return ROOT / value.removeprefix("res://")


def iter_authoritative_assets() -> list[Path]:
    paths: list[Path] = []
    art_root = ROOT / "assets" / "art"
    if art_root.exists():
        paths.extend(
            path
            for path in art_root.rglob("*")
            if path.is_file() and path.suffix.lower() in VISUAL_EXTENSIONS
        )
    fonts_root = ROOT / "assets" / "fonts"
    if fonts_root.exists():
        paths.extend(
            path
            for path in fonts_root.rglob("*")
            if path.is_file() and path.suffix.lower() in {".ttf", ".otf"}
        )
    shaders_root = ROOT / "shaders"
    if shaders_root.exists():
        paths.extend(
            path
            for path in shaders_root.rglob("*")
            if path.is_file() and path.suffix.lower() in SHADER_EXTENSIONS
        )
    return sorted(set(paths), key=lambda item: item.relative_to(ROOT).as_posix())


def sha256_for(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _svg_dimensions(path: Path) -> tuple[str, str]:
    try:
        root = ET.parse(path).getroot()
        width = root.attrib.get("width", "").removesuffix("px")
        height = root.attrib.get("height", "").removesuffix("px")
        if width and height:
            return f"{width}×{height}", "是"
        view_box = root.attrib.get("viewBox", "").split()
        if len(view_box) == 4:
            return f"{view_box[2]}×{view_box[3]}", "是"
    except (ET.ParseError, OSError):
        pass
    return "未知", "是"


def file_metadata(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    dimensions = "不适用"
    alpha = "不适用"
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            with Image.open(path) as image:
                dimensions = f"{image.width}×{image.height}"
                alpha = "是" if "A" in image.mode or "transparency" in image.info else "否"
        except OSError:
            dimensions = "无法读取"
            alpha = "未知"
    elif suffix == ".svg":
        dimensions, alpha = _svg_dimensions(path)
    elif suffix in {".ttf", ".otf"}:
        dimensions = "字体字形"
    elif suffix in SHADER_EXTENSIONS:
        dimensions = "程序生成"
    return {
        "format": suffix.removeprefix(".").upper(),
        "dimensions": dimensions,
        "alpha": alpha,
        "file_size_kb": round(path.stat().st_size / 1024, 1),
        "sha256": sha256_for(path),
    }

def classify_path(path_value: str) -> tuple[str, str, str]:
    relative = path_value.removeprefix("res://")
    parts = Path(relative).parts
    stem = Path(relative).stem.lower()
    if parts[:2] == ("assets", "fonts"):
        return "字体", "字体", "font"
    if parts and parts[0] == "shaders":
        return "宇宙", "视觉着色器", "space"

    folder = parts[2] if len(parts) > 2 and parts[:2] == ("assets", "art") else "common"
    scope, prefix = SCOPE_PREFIXES.get(folder, ("通用", "cmn"))
    lowered_parts = {part.lower() for part in parts}
    if "fonts" in lowered_parts:
        category = "字体"
    elif "buttons" in lowered_parts:
        category = "按钮"
    elif "markers" in lowered_parts:
        category = "标记与核心"
    elif "background" in lowered_parts:
        category = "背景"
    elif "effects" in lowered_parts:
        category = "特效"
    elif "ecology" in lowered_parts:
        category = "生态主体"
    elif "actors" in lowered_parts or "guardians" in lowered_parts:
        category = "角色"
    elif "obstacles" in lowered_parts:
        category = "机制障碍"
    elif "pollution_nodes" in lowered_parts:
        category = "污染节点"
    elif "stage" in lowered_parts:
        category = "棋盘舞台"
    elif "cells" in lowered_parts or (
        "board" in lowered_parts
        and any(token in stem for token in ["cell", "pentagon", "hexagon"])
    ):
        category = "棋盘格"
    elif "board" in lowered_parts:
        category = "棋盘组件"
    elif "stickers" in lowered_parts or "decorations" in lowered_parts:
        category = "装饰与生态"
    elif "notes" in lowered_parts or "pause" in lowered_parts:
        category = "纸卡与弹窗"
    elif folder == "level_select":
        category = "章节地图"
    elif folder == "tutorial":
        category = "教学界面"
    else:
        category = "界面素材"

    if stem.endswith("marker_button") or "level_marker" in stem:
        category = "地图节点"
    elif folder == "level_select" and "button" in stem:
        category = "按钮"
    return scope, category, prefix


def variant_for(path_value: str) -> str:
    stem = Path(path_value).stem.lower()
    special = {
        "facing_left": "面向左",
        "facing_right": "面向右",
        "keyboard_focus": "键盘焦点",
        "hidden_surface": "未翻开",
        "revealed_surface": "已翻开",
        "polluted_surface": "污染",
        "wrong_cloud": "误标",
        "failure_body": "失败",
        "victory_body": "胜利",
        "reaction_body": "反应",
        "light_green": "浅绿",
        "dark_green": "深绿",
    }
    for token, label in special.items():
        if stem.endswith(token) or f"_{token}_" in stem:
            return label
    for token, label in VARIANT_NAMES.items():
        if stem.endswith(f"_{token}") or stem == token:
            return label
    match = re.search(r"(?:^|_)0*(\d+)$", stem)
    if match:
        return f"变体 {int(match.group(1))}"
    return "基础"


def prompt_for(path_value: str, category: str, scope: str) -> str:
    value = path_value.lower()
    stem = Path(value).stem
    if category == "字体":
        return "PRM_FONT"
    if "shader" in value:
        return "PRM_BG_SPACE"
    if scope == "章节地图" and category not in {"按钮", "地图节点"}:
        return "PRM_MAP"
    if category == "按钮":
        return "PRM_BUTTON"
    if category in {"标记与核心", "地图节点"}:
        return "PRM_MAP_NODE" if category == "地图节点" else "PRM_MARKER"
    if category in {"纸卡与弹窗", "教学界面"}:
        if "robot" in stem or "avatar" in stem:
            return "PRM_ROBOT"
        if "button" in stem:
            return "PRM_BUTTON"
        return "PRM_PANEL"
    if category == "特效":
        return "PRM_EFFECT"
    if category == "背景":
        if scope == "欢迎页":
            return "PRM_MENU_BG"
        if scope == "陆地":
            return "PRM_BG_LAND"
        if scope == "海洋":
            return "PRM_BG_OCEAN"
        if scope == "天空":
            return "PRM_BG_SKY"
        return "PRM_STYLE"
    if category in {"棋盘舞台", "棋盘组件"}:
        return {
            "陆地": "PRM_STAGE_LAND",
            "海洋": "PRM_STAGE_OCEAN",
            "天空": "PRM_STAGE_SKY",
            "宇宙": "PRM_HUD_SPACE",
        }.get(scope, "PRM_STYLE_BOARD")
    if category == "棋盘格":
        return {
            "陆地": "PRM_CELL_LAND",
            "海洋": "PRM_CELL_OCEAN",
            "天空": "PRM_CELL_SKY",
            "宇宙": "PRM_NODE_SPACE",
        }.get(scope, "PRM_STYLE_BOARD")
    if category in {"机制障碍", "污染节点"}:
        return {
            "陆地": "PRM_MECH_LAND",
            "海洋": "PRM_MECH_OCEAN",
            "天空": "PRM_MECH_SKY",
        }.get(scope, "PRM_STYLE")
    if category == "章节地图":
        return "PRM_MAP"
    if category in {"生态主体", "装饰与生态"}:
        return "PRM_ECOLOGY"
    if category == "角色":
        if any(token in stem for token in ["slime", "monster", "pollution_core"]):
            return "PRM_MONSTER"
        if "robot" in stem:
            return "PRM_ROBOT"
        return "PRM_ECOLOGY"
    return "PRM_STYLE"


def make_short_name(path_value: str, occupied: set[str]) -> str:
    scope, _category, prefix = classify_path(path_value)
    del scope
    stem = Path(path_value).stem.lower()
    stem = re.sub(r"[^a-z0-9]+", "_", stem).strip("_")
    stem = re.sub(r"level_0*(\d+)", lambda match: f"l{int(match.group(1)):02d}", stem)
    if prefix == "land" and stem.startswith("unified_land_"):
        stem = stem.removeprefix("unified_land_")
    if prefix == "land" and stem.startswith("right_panel_guardian_"):
        stem = "right_guard_" + stem.removeprefix("right_panel_guardian_")
    for redundant in ["ocean_", "sky_", "land_", "welcome_"]:
        if stem.startswith(redundant) and prefix == redundant.removesuffix("_"):
            stem = stem[len(redundant):]
            break
    candidate = f"{prefix}_{stem}"
    if len(candidate) > MAX_SHORT_NAME_LENGTH:
        replacements = {
            "background": "bg",
            "keyboard": "key",
            "regenerate": "regen",
            "pollution": "poll",
            "guardian": "guard",
            "reaction": "react",
            "secondary": "sec",
            "primary": "pri",
            "shoulder": "shldr",
            "particles": "parts",
            "landscape": "landscape",
            "ecology": "eco",
            "surface": "surf",
            "contents": "content",
            "button": "btn",
            "marker": "mark",
        }
        tokens = candidate.split("_")
        candidate = "_".join(replacements.get(token, token) for token in tokens)
    if len(candidate) > MAX_SHORT_NAME_LENGTH:
        digest = hashlib.sha1(path_value.encode("utf-8")).hexdigest()[:4]
        candidate = candidate[: MAX_SHORT_NAME_LENGTH - 5].rstrip("_") + "_" + digest
    base = candidate
    counter = 2
    while candidate in occupied:
        suffix = f"_{counter}"
        candidate = base[: MAX_SHORT_NAME_LENGTH - len(suffix)].rstrip("_") + suffix
        counter += 1
    occupied.add(candidate)
    return candidate


def _split_asset_state(stem: str) -> tuple[str, str]:
    for suffix, label in ASSET_STATE_SUFFIXES:
        token = "_" + suffix
        if stem.endswith(token):
            return stem[: -len(token)], label
    numbered = re.fullmatch(r"(.+)_(0[1-9]|[1-9]\d*)", stem)
    if numbered and any(
        token in numbered.group(1)
        for token in ["sticker", "strip", "wall", "web"]
    ):
        return numbered.group(1), f"变体{int(numbered.group(2))}"
    return stem, ""


def asset_name_for(path_value: str, scope: str, category: str) -> str:
    stem = Path(path_value).stem.lower()
    base, state = _split_asset_state(stem)
    subject = ASSET_SUBJECT_NAMES.get(base, "")

    level_match = re.fullmatch(r"level_0*(\d+)_(.+)", base)
    if not subject and level_match:
        level = int(level_match.group(1))
        content = level_match.group(2)
        if scope == "陆地" and content in LAND_LEVEL_SUBJECTS:
            subject = f"陆地第{level}关{LAND_LEVEL_SUBJECTS[content]}"
        elif scope == "海洋" and content in OCEAN_LEVEL_SUBJECTS:
            subject = f"海洋{OCEAN_LEVEL_SUBJECTS[content]}"
        elif scope == "天空" and content == "ecology" and level in SKY_LEVEL_SUBJECTS:
            subject = f"天空第{level}关{SKY_LEVEL_SUBJECTS[level]}"

    if not subject:
        readable = base.replace("_", " ")
        subject = f"{scope}{category}：{readable}"
    if state:
        return f"{subject}（{state}）"
    return subject


def default_description(path_value: str, scope: str, category: str) -> str:
    stem = Path(path_value).stem.replace("_", " ")
    return f"{scope} · {category} · {stem}"


def parse_art_catalog() -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    aliases_by_path: dict[str, list[str]] = defaultdict(list)
    composites: dict[str, list[str]] = {}
    if not ART_CATALOG.exists():
        return aliases_by_path, composites
    text = ART_CATALOG.read_text(encoding="utf-8")
    path_pattern = re.compile(
        r"const\s+([A-Z][A-Z0-9_]*)\s*:=\s*(?:\\\s*)?"
        r'"(res://[^"\n]+)"',
        re.MULTILINE,
    )
    for alias, path_value in path_pattern.findall(text):
        aliases_by_path[path_value].append(alias)
    array_pattern = re.compile(
        r"const\s+([A-Z][A-Z0-9_]*)\s*:=\s*\[(.*?)\]",
        re.DOTALL,
    )
    known_aliases = {alias for aliases in aliases_by_path.values() for alias in aliases}
    for name, body in array_pattern.findall(text):
        dependencies = sorted(set(re.findall(r"\b[A-Z][A-Z0-9_]*\b", body)) & known_aliases)
        if dependencies:
            composites[name] = dependencies
    return aliases_by_path, composites


def iter_runtime_sources() -> list[Path]:
    paths: set[Path] = set()
    for root_name in RUNTIME_ROOTS:
        base = ROOT / root_name
        if not base.exists():
            continue
        paths.update(
            path
            for path in base.rglob("*")
            if path.is_file() and path.suffix.lower() in SOURCE_EXTENSIONS
        )
    if (ROOT / "project.godot").exists():
        paths.add(ROOT / "project.godot")
    paths.discard(ART_CATALOG)
    paths.discard(GENERATED_GDSCRIPT)
    return sorted(paths, key=lambda item: item.relative_to(ROOT).as_posix())


def collect_runtime_usage(
    aliases_by_path: dict[str, list[str]],
    composites: dict[str, list[str]],
) -> dict[str, list[str]]:
    references: dict[str, list[str]] = defaultdict(list)
    alias_to_paths: dict[str, list[str]] = defaultdict(list)
    for path_value, aliases in aliases_by_path.items():
        for alias in aliases:
            alias_to_paths[alias].append(path_value)

    source_data: list[tuple[Path, str]] = []
    for source in iter_runtime_sources():
        try:
            source_data.append((source, source.read_text(encoding="utf-8")))
        except UnicodeDecodeError:
            continue

    for source, text in source_data:
        relative = source.relative_to(ROOT).as_posix()
        for path_value in aliases_by_path:
            start = 0
            while True:
                index = text.find(path_value, start)
                if index < 0:
                    break
                line = text.count("\n", 0, index) + 1
                location = f"{relative}:{line}"
                if location not in references[path_value]:
                    references[path_value].append(location)
                start = index + len(path_value)
        for alias, path_values in alias_to_paths.items():
            match = re.search(rf"\b{re.escape(alias)}\b", text)
            if match is None:
                continue
            line = text.count("\n", 0, match.start()) + 1
            location = f"{relative}:{line} via {alias}"
            for path_value in path_values:
                if location not in references[path_value]:
                    references[path_value].append(location)
        for composite, dependencies in composites.items():
            match = re.search(rf"\b{re.escape(composite)}\b", text)
            if match is None:
                continue
            line = text.count("\n", 0, match.start()) + 1
            location = f"{relative}:{line} via {composite}"
            for dependency in dependencies:
                for path_value in alias_to_paths.get(dependency, []):
                    if location not in references[path_value]:
                        references[path_value].append(location)

    known_paths = set(aliases_by_path)
    for source, text in source_data:
        relative = source.relative_to(ROOT).as_posix()
        for match in re.finditer(r'res://(?:assets/(?:art|fonts)|shaders)/[^"\'\s)\]]+', text):
            path_value = match.group(0)
            if path_value in known_paths and references.get(path_value):
                continue
            line = text.count("\n", 0, match.start()) + 1
            location = f"{relative}:{line}"
            if location not in references[path_value]:
                references[path_value].append(location)
    return references

def refresh_entries(entries: list[dict[str, Any]]) -> None:
    aliases_by_path, composites = parse_art_catalog()
    usage = collect_runtime_usage(aliases_by_path, composites)
    for entry in entries:
        path_value = entry["path"]
        path = local_path(path_value)
        entry["aliases"] = sorted(aliases_by_path.get(path_value, []))
        entry["sources"] = usage.get(path_value, [])
        if path.exists() and path.is_file():
            entry.update(file_metadata(path))
            if entry["sources"]:
                entry["runtime_status"] = "运行中"
            elif entry["aliases"]:
                entry["runtime_status"] = "已登记"
            else:
                entry["runtime_status"] = "未引用"
        else:
            entry["runtime_status"] = "文件缺失"


def new_entry(
    path: Path,
    asset_id: str,
    occupied_names: set[str],
) -> dict[str, Any]:
    path_value = res_path(path)
    scope, category, _prefix = classify_path(path_value)
    entry: dict[str, Any] = {
        "id": asset_id,
        "short_name": make_short_name(path_value, occupied_names),
        "asset_name": asset_name_for(path_value, scope, category),
        "description": default_description(path_value, scope, category),
        "scope": scope,
        "category": category,
        "variant": variant_for(path_value),
        "path": path_value,
        "format": "",
        "dimensions": "",
        "alpha": "",
        "file_size_kb": 0.0,
        "sha256": "",
        "runtime_status": "未引用",
        "aliases": [],
        "sources": [],
        "prompt_id": prompt_for(path_value, category, scope),
        "replacement_status": "待评估",
        "rights_status": "待复核",
        "source_tool": "",
        "commercial_evidence": "",
        "note": "",
        "review": "待复核",
    }
    entry.update(file_metadata(path))
    return entry


def load_payload() -> dict[str, Any]:
    payload = json.loads(MANIFEST.read_text(encoding="utf-8"))
    if payload.get("schema") not in {1, 2}:
        raise SystemExit(f"Unsupported art catalog schema: {payload.get('schema')}")
    for entry in payload.get("entries", []):
        if not entry.get("asset_name"):
            entry["asset_name"] = asset_name_for(
                entry["path"], entry["scope"], entry["category"]
            )
    payload["schema"] = 2
    return payload


def save_payload(payload: dict[str, Any]) -> None:
    payload["schema"] = 2
    payload["id_format"] = "ART_####"
    payload["short_name_format"] = "<scope>_<subject>_<state>"
    payload["short_name_max_length"] = MAX_SHORT_NAME_LENGTH
    payload["source_of_truth"] = "docs/art/GAME_ART_CATALOG.xlsx"
    payload["prompt_document"] = "docs/art/AI_ART_PROMPT_SYSTEM.md"
    highest_assigned = max(
        [int(entry["id"].split("_")[1]) for entry in payload["entries"]] or [0]
    )
    payload["max_id"] = max(int(payload.get("max_id", 0)), highest_assigned)
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def gd_quote(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def generated_gdscript_text(entries: list[dict[str, Any]]) -> str:
    ordered = sorted(entries, key=lambda item: int(item["id"].split("_")[1]))
    lines = [
        "# Generated by tools/art_catalog.py. Do not edit by hand.",
        "class_name ArtIds",
        "extends RefCounted",
        "",
        "const PATHS := {",
    ]
    for entry in ordered:
        lines.append(f'\t&{gd_quote(entry["id"])}: {gd_quote(entry["path"])},')
    lines.extend(["}", "", "const SHORT_NAMES := {"])
    for entry in ordered:
        lines.append(
            f'\t&{gd_quote(entry["id"])}: &{gd_quote(entry["short_name"])},'
        )
    lines.extend(["}", "", "const IDS_BY_PATH := {"])
    for entry in sorted(ordered, key=lambda item: item["path"]):
        lines.append(f'\t{gd_quote(entry["path"])}: &{gd_quote(entry["id"])},')
    lines.extend(
        [
            "}",
            "",
            "",
            "static func path_for(asset_id: StringName) -> String:",
            '\treturn str(PATHS.get(asset_id, ""))',
            "",
            "",
            "static func short_name_for(asset_id: StringName) -> StringName:",
            '\treturn StringName(SHORT_NAMES.get(asset_id, &""))',
            "",
            "",
            "static func id_for_path(path: String) -> StringName:",
            '\treturn StringName(IDS_BY_PATH.get(path, &""))',
            "",
        ]
    )
    return "\n".join(lines)


def write_generated_gdscript(entries: list[dict[str, Any]]) -> None:
    GENERATED_GDSCRIPT.write_text(generated_gdscript_text(entries), encoding="utf-8")


def style_header(sheet: Any, color: str = "31583F") -> None:
    fill = PatternFill("solid", fgColor=color)
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30


def add_list_validation(
    sheet: Any,
    column_letter: str,
    values: list[str],
    max_row: int,
) -> None:
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(values) + '"',
        allow_blank=False,
    )
    sheet.add_data_validation(validation)
    validation.add(f"{column_letter}2:{column_letter}{max(max_row, 5000)}")


def export_workbook(entries: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "美术总表"
    sheet.append(WORKBOOK_HEADERS)
    style_header(sheet)

    scope_colors = {
        "通用": "E2F0D9",
        "玩法通用": "DDEBF7",
        "教学": "FCE4D6",
        "欢迎页": "FFF2CC",
        "陆地": "E2F0D9",
        "章节地图": "FFE699",
        "海洋": "DDEBF7",
        "天空": "E4DFEC",
        "宇宙": "D9E1F2",
        "字体": "E7E6E6",
    }
    for entry in sorted(entries, key=lambda item: int(item["id"].split("_")[1])):
        sheet.append(
            [
                entry["id"],
                entry["short_name"],
                entry["asset_name"],
                entry["description"],
                entry["scope"],
                entry["category"],
                entry["variant"],
                entry["path"],
                entry.get("format", ""),
                entry.get("dimensions", ""),
                entry.get("alpha", ""),
                entry.get("file_size_kb", 0.0),
                entry.get("sha256", ""),
                entry.get("runtime_status", ""),
                "\n".join(entry.get("aliases", [])),
                "\n".join(entry.get("sources", [])),
                entry["prompt_id"],
                entry.get("replacement_status", "待评估"),
                entry.get("rights_status", "待复核"),
                entry.get("source_tool", ""),
                entry.get("commercial_evidence", ""),
                entry.get("note", ""),
                entry.get("review", "待复核"),
            ]
        )
        row = sheet.max_row
        sheet.cell(row, 5).fill = PatternFill(
            "solid", fgColor=scope_colors.get(entry["scope"], "FFFFFF")
        )
        for column in range(1, len(WORKBOOK_HEADERS) + 1):
            sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
        for column in [1, 9, 10, 11, 12, 14, 17, 18, 19, 23]:
            sheet.cell(row, column).alignment = Alignment(
                horizontal="center", vertical="top", wrap_text=True
            )

    widths = [
        13, 38, 38, 42, 13, 17, 15, 66, 11, 15, 13, 14, 18,
        14, 36, 56, 20, 15, 15, 24, 36, 36, 14,
    ]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "H2"
    sheet.auto_filter.ref = sheet.dimensions

    prompt_validation = DataValidation(
        type="list",
        formula1=f"'提示词索引'!$A$2:$A${len(PROMPTS) + 1}",
        allow_blank=False,
    )
    sheet.add_data_validation(prompt_validation)
    prompt_validation.add(f"Q2:Q{max(sheet.max_row, 5000)}")
    add_list_validation(sheet, "R", ["待评估", "保留", "待重制", "重制中", "已确认"], sheet.max_row)
    add_list_validation(sheet, "S", ["待复核", "已确认商用", "需替换"], sheet.max_row)
    add_list_validation(sheet, "W", ["待复核", "修改中", "已确认"], sheet.max_row)
    sheet.conditional_formatting.add(
        f"W2:W{sheet.max_row}",
        FormulaRule(formula=['$W2="已确认"'], fill=PatternFill("solid", fgColor="C6E0B4")),
    )
    sheet.conditional_formatting.add(
        f"S2:S{sheet.max_row}",
        FormulaRule(formula=['$S2="需替换"'], fill=PatternFill("solid", fgColor="F4CCCC")),
    )

    naming = workbook.create_sheet("命名规则")
    naming.append(["项目", "规则", "示例", "说明"])
    naming_rows = [
        ["永久代号", "ART_####", "ART_0001", "自动分配；删除或淘汰后也禁止复用"],
        ["中文素材名", "章节 + 具体对象 + 状态", "海洋暂停按钮（悬停）", "必须一眼看出是什么素材、用于哪里；不能只写按钮、背景或角色"],
        ["规范短名", "<scope>_<subject>_<state>", "ocean_whale_fall_healthy", "小写 snake_case，最多36字符；可人工精修但不得重复"],
        ["范围前缀", "cmn/play/tut/wel/land/map/ocean/sky/space/font", "sky_hexagon_hidden_surface", "先写使用范围，再写内容和状态"],
        ["物理文件名", "<location>_<role>_<content>_<state>.<ext>", "menu_secondary_pressed.svg", "文件移动或改名时更新原代号路径，不另建代号"],
        ["新增素材", "只追加新代号", "ART_0211", "运行 discover；不得重排旧记录"],
        ["缺失素材", "保留原记录", "运行状态=文件缺失", "用于追踪移动、误删和历史替换关系"],
    ]
    for row in naming_rows:
        naming.append(row)
    style_header(naming)
    for index, width in enumerate([18, 44, 42, 68], 1):
        naming.column_dimensions[get_column_letter(index)].width = width
    for row in naming.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    naming.freeze_panes = "A2"

    prompt_sheet = workbook.create_sheet("提示词索引")
    prompt_sheet.append(["提示词代号", "名称", "文档章节", "适用内容", "提示词文档"])
    for prompt_id, name, section, usage in PROMPTS:
        prompt_sheet.append([prompt_id, name, section, usage, "AI_ART_PROMPT_SYSTEM.md"])
        link_cell = prompt_sheet.cell(prompt_sheet.max_row, 5)
        link_cell.hyperlink = "AI_ART_PROMPT_SYSTEM.md"
        link_cell.style = "Hyperlink"
    style_header(prompt_sheet, "385D78")
    for index, width in enumerate([23, 26, 25, 50, 34], 1):
        prompt_sheet.column_dimensions[get_column_letter(index)].width = width
    for row in prompt_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    prompt_sheet.freeze_panes = "A2"
    prompt_sheet.auto_filter.ref = prompt_sheet.dimensions

    enums = workbook.create_sheet("枚举与状态")
    enums.append(["字段", "可选值", "含义"])
    enum_rows = [
        ["运行状态", "运行中", "程序或场景通过路径/ArtCatalog别名引用"],
        ["运行状态", "已登记", "已进入ArtCatalog，但当前未发现外部使用位置"],
        ["运行状态", "未引用", "文件存在，但尚未发现运行引用"],
        ["运行状态", "文件缺失", "永久记录仍保留，但当前路径不存在"],
        ["重制状态", "待评估 / 保留 / 待重制 / 重制中 / 已确认", "跟踪正式美术替换进度"],
        ["权利状态", "待复核 / 已确认商用 / 需替换", "权利确认必须有凭证支持"],
        ["审核状态", "待复核 / 修改中 / 已确认", "目录字段人工复核状态，不代表视觉效果自动通过"],
    ]
    for row in enum_rows:
        enums.append(row)
    style_header(enums, "74643C")
    for index, width in enumerate([18, 48, 70], 1):
        enums.column_dimensions[get_column_letter(index)].width = width
    for row in enums.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    guide = workbook.create_sheet("使用说明")
    guide.column_dimensions["A"].width = 115
    instructions = [
        "《绿色扫雷》美术素材维护说明",
        "1. “代号”是永久内部键，不得修改、重排或复用；文件路径和短名改变时仍保留原代号。",
        "2. “素材名称（中文）”必须直接写清具体对象与状态，例如“海洋暂停按钮（悬停）”，不能只写“按钮”；“规范短名”采用小写 snake_case，最长36字符且不得重复。",
        "3. “当前路径”可在文件移动/改名后人工更新；必须是 res:// 开头的正式资源路径。",
        "4. 尺寸、透明通道、大小、SHA-256、运行状态、程序别名和使用位置由工具计算，不要手工维护。",
        "5. “提示词代号”关联 docs/art/AI_ART_PROMPT_SYSTEM.md；重制时先查“提示词索引”，并同时使用 PRM_STYLE 全局母提示词。",
        "6. 来源/生成工具、商用凭证和权利状态必须按真实证据填写；“已确认商用”不应只凭主观判断。",
        "7. Excel修改后运行：python tools/art_catalog.py import。导入按代号同步，不依赖文件排序或源码行号。",
        "8. 放入新素材后运行：python tools/art_catalog.py discover。工具只追加新ID；能唯一匹配原SHA-256的移动文件会保留旧ID。",
        "9. 重新生成Excel运行：python tools/art_catalog.py export。导出前先import尚未同步的Excel修改，避免覆盖。",
        "10. 提交前运行：python tools/art_catalog.py validate。该命令只校验结构和一致性，不判断画面质量。",
        "11. Excel无法保存图案本身；实际图案始终由“代号 + 当前路径 + SHA-256”定位。",
    ]
    for index, text in enumerate(instructions, 1):
        guide.cell(index, 1, text)
        guide.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="top")
    guide.cell(1, 1).font = Font(size=16, bold=True, color="31583F")

    metadata = workbook.create_sheet("元数据")
    metadata.sheet_state = "hidden"
    metadata.append(["schema", 2])
    metadata.append(["id_format", "ART_####"])
    metadata.append(["entry_count", len(entries)])
    metadata.append(["short_name_max_length", MAX_SHORT_NAME_LENGTH])
    metadata.append(["source_of_truth", "docs/art/GAME_ART_CATALOG.xlsx"])
    metadata.append(["prompt_document", "docs/art/AI_ART_PROMPT_SYSTEM.md"])

    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(WORKBOOK)


def initialize() -> None:
    if MANIFEST.exists():
        raise SystemExit(f"Refusing to initialize: {MANIFEST} already exists")
    occupied: set[str] = set()
    entries = [
        new_entry(path, f"ART_{index:04d}", occupied)
        for index, path in enumerate(iter_authoritative_assets(), 1)
    ]
    payload = {"schema": 2, "max_id": len(entries), "entries": entries}
    refresh_entries(entries)
    save_payload(payload)
    write_generated_gdscript(entries)
    export_workbook(entries)
    print(f"Initialized {len(entries)} stable art IDs")
    print(f"Workbook: {WORKBOOK}")


def discover() -> None:
    payload = load_payload()
    entries = payload["entries"]
    current = {res_path(path): path for path in iter_authoritative_assets()}
    by_path = {entry["path"]: entry for entry in entries}
    missing = [entry for entry in entries if entry["path"] not in current]
    untracked = {path_value: path for path_value, path in current.items() if path_value not in by_path}

    missing_by_hash: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for entry in missing:
        if entry.get("sha256"):
            missing_by_hash[entry["sha256"]].append(entry)
    untracked_by_hash: dict[str, list[tuple[str, Path]]] = defaultdict(list)
    for path_value, path in untracked.items():
        untracked_by_hash[sha256_for(path)].append((path_value, path))

    moved = 0
    for checksum, old_entries in missing_by_hash.items():
        new_paths = untracked_by_hash.get(checksum, [])
        if len(old_entries) != 1 or len(new_paths) != 1:
            continue
        entry = old_entries[0]
        new_path_value, _new_path = new_paths[0]
        old_path_value = entry["path"]
        entry["path"] = new_path_value
        untracked.pop(new_path_value, None)
        moved += 1
        print(f"Preserved {entry['id']} for moved asset: {old_path_value} -> {new_path_value}")

    occupied = {entry["short_name"] for entry in entries}
    next_id = max(int(payload.get("max_id", 0)), max(
        [int(entry["id"].split("_")[1]) for entry in entries] or [0]
    ))
    added = 0
    for path_value in sorted(untracked):
        next_id += 1
        entries.append(new_entry(untracked[path_value], f"ART_{next_id:04d}", occupied))
        added += 1

    payload["max_id"] = next_id
    entries.sort(key=lambda item: int(item["id"].split("_")[1]))
    refresh_entries(entries)
    save_payload(payload)
    write_generated_gdscript(entries)
    export_workbook(entries)
    print(f"Discovered {added} new assets, preserved {moved} moved assets")
    print(f"Catalog now contains {len(entries)} permanent IDs")


def import_workbook() -> None:
    payload = load_payload()
    entries = payload["entries"]
    by_id = {entry["id"]: entry for entry in entries}
    workbook = load_workbook(WORKBOOK)
    if "美术总表" not in workbook.sheetnames:
        raise SystemExit("Workbook is missing sheet: 美术总表")
    sheet = workbook["美术总表"]
    headers = {str(cell.value or "").strip(): index for index, cell in enumerate(sheet[1], 1)}
    required = {"代号", *EDITABLE_COLUMNS}
    missing_headers = required - set(headers)
    if missing_headers:
        raise SystemExit(f"Workbook is missing columns: {sorted(missing_headers)}")

    seen: set[str] = set()
    for row in range(2, sheet.max_row + 1):
        asset_id = str(sheet.cell(row, headers["代号"]).value or "").strip()
        if not asset_id:
            continue
        if asset_id in seen:
            raise SystemExit(f"Duplicate art ID in workbook: {asset_id}")
        if asset_id not in by_id:
            raise SystemExit(
                f"Unknown art ID in workbook: {asset_id}. Add files with the discover command."
            )
        seen.add(asset_id)
        entry = by_id[asset_id]
        for column, field in EDITABLE_COLUMNS.items():
            value = str(sheet.cell(row, headers[column]).value or "").strip()
            if field == "path":
                value = value.replace("\\", "/")
            entry[field] = value

    missing_ids = set(by_id) - seen
    if missing_ids:
        raise SystemExit(f"Workbook is missing art IDs: {sorted(missing_ids)}")
    errors = basic_errors(payload)
    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)

    refresh_entries(entries)
    save_payload(payload)
    write_generated_gdscript(entries)
    export_workbook(entries)
    print(f"Imported {len(entries)} rows and rebuilt JSON, Excel, and Godot lookup data")

def basic_errors(payload: dict[str, Any]) -> list[str]:
    entries = payload.get("entries", [])
    errors: list[str] = []
    ids = [str(entry.get("id", "")) for entry in entries]
    names = [str(entry.get("short_name", "")) for entry in entries]
    paths = [str(entry.get("path", "")) for entry in entries]
    if len(ids) != len(set(ids)):
        errors.append("Manifest contains duplicate art IDs")
    if len(names) != len(set(names)):
        errors.append("Manifest contains duplicate short names")
    if len(paths) != len(set(paths)):
        errors.append("Manifest contains duplicate asset paths")

    numeric_ids: list[int] = []
    for entry in entries:
        asset_id = str(entry.get("id", ""))
        short_name = str(entry.get("short_name", ""))
        asset_name = str(entry.get("asset_name", "")).strip()
        path_value = str(entry.get("path", ""))
        if not ID_PATTERN.fullmatch(asset_id):
            errors.append(f"Invalid ID format: {asset_id}")
        else:
            numeric_ids.append(int(asset_id.split("_")[1]))
        if not SHORT_NAME_PATTERN.fullmatch(short_name):
            errors.append(f"Invalid short name: {asset_id} {short_name!r}")
        if len(short_name) > MAX_SHORT_NAME_LENGTH:
            errors.append(
                f"Short name exceeds {MAX_SHORT_NAME_LENGTH} characters: {asset_id} {short_name}"
            )
        if not asset_name:
            errors.append(f"Chinese asset name is empty: {asset_id}")
        if not path_value.startswith("res://"):
            errors.append(f"Path does not start with res://: {asset_id} {path_value}")
        else:
            suffix = Path(path_value).suffix.lower()
            valid_extension = suffix in VISUAL_EXTENSIONS | SHADER_EXTENSIONS
            valid_root = path_value.startswith(
                ("res://assets/art/", "res://assets/fonts/", "res://shaders/")
            )
            if not valid_extension or not valid_root:
                errors.append(f"Path is outside authoritative art roots: {asset_id} {path_value}")
        if entry.get("prompt_id") not in PROMPT_IDS:
            errors.append(f"Unknown prompt ID: {asset_id} {entry.get('prompt_id')}")
        if entry.get("replacement_status") not in {
            "待评估", "保留", "待重制", "重制中", "已确认"
        }:
            errors.append(f"Invalid replacement status: {asset_id}")
        if entry.get("rights_status") not in {"待复核", "已确认商用", "需替换"}:
            errors.append(f"Invalid rights status: {asset_id}")
        if entry.get("review") not in {"待复核", "修改中", "已确认"}:
            errors.append(f"Invalid review status: {asset_id}")
    if numeric_ids and int(payload.get("max_id", 0)) < max(numeric_ids):
        errors.append("max_id is lower than an assigned ID; old IDs could be reused")
    return errors


def workbook_errors(entries: list[dict[str, Any]]) -> list[str]:
    if not WORKBOOK.exists():
        return [f"Workbook is missing: {WORKBOOK}"]
    errors: list[str] = []
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=False)
    if "美术总表" not in workbook.sheetnames:
        return ["Workbook is missing sheet: 美术总表"]
    sheet = workbook["美术总表"]
    headers = {str(cell.value or "").strip(): index for index, cell in enumerate(sheet[1], 1)}
    required = {"代号", *EDITABLE_COLUMNS}
    missing_headers = required - set(headers)
    if missing_headers:
        return [f"Workbook is missing columns: {sorted(missing_headers)}"]
    rows: dict[str, dict[str, str]] = {}
    for row in range(2, sheet.max_row + 1):
        asset_id = str(sheet.cell(row, headers["代号"]).value or "").strip()
        if not asset_id:
            continue
        if asset_id in rows:
            errors.append(f"Duplicate art ID in workbook: {asset_id}")
            continue
        rows[asset_id] = {
            field: str(sheet.cell(row, headers[column]).value or "").strip().replace("\\", "/")
            if field == "path"
            else str(sheet.cell(row, headers[column]).value or "").strip()
            for column, field in EDITABLE_COLUMNS.items()
        }
    expected = {entry["id"]: entry for entry in entries}
    if set(rows) != set(expected):
        missing = sorted(set(expected) - set(rows))
        extra = sorted(set(rows) - set(expected))
        if missing:
            errors.append(f"Workbook is missing IDs: {missing}")
        if extra:
            errors.append(f"Workbook has unknown IDs: {extra}")
    for asset_id in sorted(set(rows) & set(expected)):
        for field in EDITABLE_COLUMNS.values():
            if rows[asset_id][field] != str(expected[asset_id].get(field, "")):
                errors.append(f"Workbook differs from manifest: {asset_id} {field}")
    return errors


def validate() -> None:
    payload = load_payload()
    entries = payload["entries"]
    errors = basic_errors(payload)
    manifest_paths = {entry["path"] for entry in entries}
    current_paths = {res_path(path) for path in iter_authoritative_assets()}
    untracked = sorted(current_paths - manifest_paths)
    if untracked:
        errors.append(f"Authoritative files lack IDs: {untracked}")

    aliases_by_path, _composites = parse_art_catalog()
    unknown_alias_paths = sorted(set(aliases_by_path) - manifest_paths)
    if unknown_alias_paths:
        errors.append(f"ArtCatalog paths lack stable IDs: {unknown_alias_paths}")

    refreshed = json.loads(json.dumps(entries, ensure_ascii=False))
    refresh_entries(refreshed)
    computed_fields = [
        "format", "dimensions", "alpha", "file_size_kb", "sha256",
        "runtime_status", "aliases", "sources",
    ]
    refreshed_by_id = {entry["id"]: entry for entry in refreshed}
    for entry in entries:
        current = refreshed_by_id[entry["id"]]
        for field in computed_fields:
            if entry.get(field) != current.get(field):
                errors.append(f"Computed metadata is stale: {entry['id']} {field}")

    known_ids = {entry["id"] for entry in entries}
    for source in iter_runtime_sources():
        text = source.read_text(encoding="utf-8")
        for asset_id in set(re.findall(r"\bART_\d{4}\b", text)):
            if asset_id not in known_ids:
                errors.append(
                    f"Unknown stable art ID in {source.relative_to(ROOT).as_posix()}: {asset_id}"
                )

    expected_generated = generated_gdscript_text(entries)
    if not GENERATED_GDSCRIPT.exists():
        errors.append(f"Generated lookup file is missing: {GENERATED_GDSCRIPT}")
    elif GENERATED_GDSCRIPT.read_text(encoding="utf-8") != expected_generated:
        errors.append("scripts/art_ids.gd differs from the manifest")

    errors.extend(workbook_errors(entries))
    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
    missing_count = sum(entry.get("runtime_status") == "文件缺失" for entry in entries)
    print(
        f"Validated {len(entries)} stable art IDs, metadata, workbook, prompts, "
        f"Godot lookups, and source references ({missing_count} missing historical files)"
    )


def export_catalog() -> None:
    payload = load_payload()
    entries = payload["entries"]
    refresh_entries(entries)
    save_payload(payload)
    write_generated_gdscript(entries)
    export_workbook(entries)
    print(f"Exported {len(entries)} rows to {WORKBOOK}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "command",
        choices=["initialize", "discover", "export", "import", "validate"],
    )
    args = parser.parse_args()
    if args.command == "initialize":
        initialize()
    elif args.command == "discover":
        discover()
    elif args.command == "export":
        export_catalog()
    elif args.command == "import":
        import_workbook()
    else:
        validate()


if __name__ == "__main__":
    main()
