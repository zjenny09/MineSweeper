#!/usr/bin/env python3
"""Maintain stable in-game copy IDs, PO catalogs, and the review workbook.

Commands:
  initialize  One-time migration from Chinese msgids to TXT_#### IDs.
  export      Rebuild the Excel workbook from copy_catalog.json.
  import      Apply edited Excel text to JSON and both PO catalogs.
  validate    Check IDs, catalogs, placeholders, and source usage.
"""

from __future__ import annotations

import argparse
import ast
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
EN_PO = ROOT / "localization" / "green_sweeper.en.po"
ZH_PO = ROOT / "localization" / "green_sweeper.zh_CN.po"
MANIFEST = ROOT / "localization" / "copy_catalog.json"
WORKBOOK = ROOT / "docs" / "copy" / "GAME_COPY_CATALOG.xlsx"
ID_PATTERN = re.compile(r"^TXT_\d{4}$")
QUOTED_STRING_PATTERN = re.compile(r'"(?:\\.|[^"\\])*"')
PLACEHOLDER_PATTERN = re.compile(r"%(?:\d+\$)?[-+0 #]*(?:\d+|\*)?(?:\.\d+)?[diouxXeEfFgGcs]")
HAN_PATTERN = re.compile(r"[一-龥]")
SOURCE_EXTENSIONS = {".gd", ".tscn", ".godot"}
SOURCE_EXCLUDES = {
    Path("scripts/audio_manager.gd"),
}

OLD_COPY = {
    "第15关 · 云冠（球面原型）",
    "旋转足球烯球面，净化全部安全面",
    "拖拽旋转 · 左键净化 · 右键标记\n",
}
DEV_COPY = {
    "操作方法\n左键净化 · 右键标记\n双击数字快速展开",
    "净化时间\n00:00.0\n\n污染核心    0/%d\n\n净化状态\n准备中",
    "净化目标\n清除所有污染核心\n恢复%s生态\n\n",
    "第%d关 · %s\n%d × %d 棋盘\n污染核心 %d\n\n",
    "第1关:1",
    "第2关:2",
    "第3关:3",
    "第4关:4",
    "第5关:5",
    "无效启动参数；请使用 --level=1 至 --level=21 或 --universe。",
}
NON_RENDER_COPY = {"退出", "取消", "返回", "海洋篇"}


def _po_value(token: str) -> str:
    return ast.literal_eval(token.strip())


def parse_po(path: Path) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    current: dict[str, Any] | None = None
    active_field = ""

    for line_number, line in enumerate(lines, 1):
        if line.startswith("msgid "):
            if current is not None and current.get("msgid"):
                entries.append(current)
            current = {
                "msgid": _po_value(line[6:]),
                "msgstr": "",
                "line": line_number,
            }
            active_field = "msgid"
        elif line.startswith("msgstr ") and current is not None:
            current["msgstr"] = _po_value(line[7:])
            active_field = "msgstr"
        elif line.startswith('"') and current is not None and active_field:
            current[active_field] += _po_value(line)
        elif not line.strip():
            if current is not None and current.get("msgid"):
                entries.append(current)
            current = None
            active_field = ""

    if current is not None and current.get("msgid"):
        entries.append(current)
    return entries


def po_quote(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def write_po(path: Path, locale: str, entries: list[dict[str, Any]]) -> None:
    header = [
        'msgid ""',
        'msgstr ""',
        f'"Project-Id-Version: Green Sweeper 0.7.0\\n"',
        f'"Language: {locale}\\n"',
        '"MIME-Version: 1.0\\n"',
        '"Content-Type: text/plain; charset=UTF-8\\n"',
        '"Content-Transfer-Encoding: 8bit\\n"',
        "",
    ]
    chunks = ["\n".join(header)]
    language_key = "zh_CN" if locale == "zh_CN" else "en"
    for entry in entries:
        chunks.append(
            "\n".join(
                [
                    f"#. {entry['category']} | {entry['status']}",
                    f"msgid {po_quote(entry['id'])}",
                    f"msgstr {po_quote(entry[language_key])}",
                    "",
                ]
            )
        )
    path.write_text("\n".join(chunks).rstrip() + "\n", encoding="utf-8")


def placeholders(value: str) -> list[str]:
    return PLACEHOLDER_PATTERN.findall(value)


def category_for_line(line: int, text: str) -> str:
    if text in DEV_COPY:
        return "开发与预览"
    if text in OLD_COPY:
        return "旧版原型"
    if 9 <= line <= 43:
        return "品牌与主菜单"
    if 45 <= line <= 100:
        return "设置"
    if 102 <= line <= 223:
        return "地图与导航"
    if 225 <= line <= 341:
        return "关卡名称与棋盘"
    if 342 <= line <= 427:
        return "目标与游戏状态"
    if 429 <= line <= 475:
        return "暂停、失败与结算"
    if 477 <= line <= 557:
        return "操作与章节规则"
    if 558 <= line <= 683:
        return "扫描与格子提示"
    if 684 <= line <= 727:
        return "天空动态状态"
    if 729 <= line <= 763:
        return "引导通用"
    if 765 <= line <= 901:
        return "陆地与海洋引导"
    if 903 <= line <= 974:
        return "天空引导"
    if 975 <= line <= 1018:
        return "宇宙引导"
    if 1020 <= line <= 1153:
        return "宇宙篇"
    if "漂流" in text or "洋流" in text or "潮区" in text or "过滤边界" in text:
        return "海洋动态状态"
    if any(token in text for token in ["按 J", "按 K", "白色=", "数字格", "检测种子", "快速展开"]):
        return "第一关教学"
    return "补充与兼容文案"


def status_for_text(text: str) -> str:
    if text in OLD_COPY:
        return "旧版"
    if text in DEV_COPY:
        return "开发"
    if text in NON_RENDER_COPY:
        return "非独立显示"
    return "运行"


def iter_source_files() -> list[Path]:
    files: list[Path] = []
    for base in [ROOT / "scripts", ROOT / "scenes"]:
        for path in base.rglob("*"):
            if path.is_file() and path.suffix in SOURCE_EXTENSIONS:
                relative = path.relative_to(ROOT)
                if relative not in SOURCE_EXCLUDES:
                    files.append(path)
    return sorted(files)


def decode_quoted(token: str) -> str | None:
    try:
        return json.loads(token)
    except (json.JSONDecodeError, TypeError):
        try:
            return ast.literal_eval(token)
        except (SyntaxError, ValueError):
            return None


def collect_occurrences(values: set[str]) -> dict[str, list[str]]:
    result: dict[str, list[str]] = defaultdict(list)
    for path in iter_source_files():
        relative = path.relative_to(ROOT).as_posix()
        text = path.read_text(encoding="utf-8")
        for match in QUOTED_STRING_PATTERN.finditer(text):
            value = decode_quoted(match.group(0))
            if value not in values:
                continue
            line = text.count("\n", 0, match.start()) + 1
            location = f"{relative}:{line}"
            if location not in result[value]:
                result[value].append(location)
    return result


def refresh_source_locations(entries: list[dict[str, Any]]) -> None:
    by_id = {entry["id"]: entry for entry in entries}
    locations: dict[str, list[str]] = defaultdict(list)
    for path in iter_source_files():
        relative = path.relative_to(ROOT).as_posix()
        text = path.read_text(encoding="utf-8")
        for match in QUOTED_STRING_PATTERN.finditer(text):
            value = decode_quoted(match.group(0))
            if not isinstance(value, str) or value not in by_id:
                continue
            line = text.count("\n", 0, match.start()) + 1
            location = f"{relative}:{line}"
            if location not in locations[value]:
                locations[value].append(location)
    for entry in entries:
        entry["sources"] = locations.get(entry["id"], [])


def replace_source_literals(mapping: dict[str, str]) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    for path in iter_source_files():
        text = path.read_text(encoding="utf-8")

        def replace(match: re.Match[str]) -> str:
            value = decode_quoted(match.group(0))
            copy_id = mapping.get(value or "")
            if copy_id is None:
                return match.group(0)
            counts[copy_id] += 1
            return json.dumps(copy_id)

        updated = QUOTED_STRING_PATTERN.sub(replace, text)
        if updated != text:
            path.write_text(updated, encoding="utf-8")
    return counts


def save_manifest(entries: list[dict[str, Any]]) -> None:
    payload = {
        "schema": 1,
        "id_format": "TXT_####",
        "source_of_truth": "docs/copy/GAME_COPY_CATALOG.xlsx",
        "entries": entries,
    }
    MANIFEST.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def load_manifest() -> list[dict[str, Any]]:
    payload = json.loads(MANIFEST.read_text(encoding="utf-8"))
    return payload["entries"]


def export_workbook(entries: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文案总表"
    headers = [
        "代号",
        "分类",
        "状态",
        "中文",
        "English",
        "占位符",
        "使用位置（仅供参考）",
        "修改备注",
        "审核状态",
    ]
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="31583F")
    header_font = Font(color="FFFFFF", bold=True)
    category_colors = {
        "品牌与主菜单": "DDEED9",
        "设置": "E8E2F3",
        "地图与导航": "FFF0C2",
        "关卡名称与棋盘": "E3F0D2",
        "目标与游戏状态": "F6E3C8",
        "暂停、失败与结算": "F4D6D6",
        "操作与章节规则": "D9EBF4",
        "扫描与格子提示": "DDEBF7",
        "天空动态状态": "D8F2F3",
        "引导通用": "F6E5F0",
        "陆地与海洋引导": "E1EFDA",
        "天空引导": "D7EEF7",
        "宇宙引导": "DDD9F0",
        "宇宙篇": "D4D7EB",
        "海洋动态状态": "D4EEF2",
        "第一关教学": "FCE4D6",
        "旧版原型": "E7E6E6",
        "开发与预览": "E7E6E6",
        "补充与兼容文案": "F2F2F2",
    }

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for entry in entries:
        sheet.append(
            [
                entry["id"],
                entry["category"],
                entry["status"],
                entry["zh_CN"],
                entry["en"],
                "、".join(entry["placeholders"]),
                "\n".join(entry.get("sources", [])),
                entry.get("note", ""),
                entry.get("review", "待复核"),
            ]
        )
        row = sheet.max_row
        fill = PatternFill(
            "solid",
            fgColor=category_colors.get(entry["category"], "FFFFFF"),
        )
        sheet.cell(row, 2).fill = fill
        for column in range(1, len(headers) + 1):
            sheet.cell(row, column).alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    widths = [13, 20, 13, 48, 56, 22, 52, 36, 14]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 28

    review_validation = DataValidation(
        type="list",
        formula1='"待复核,修改中,已确认"',
        allow_blank=False,
    )
    sheet.add_data_validation(review_validation)
    review_validation.add(f"I2:I{sheet.max_row}")
    confirmed_fill = PatternFill("solid", fgColor="C6E0B4")
    sheet.conditional_formatting.add(
        f"I2:I{sheet.max_row}",
        FormulaRule(formula=['$I2="已确认"'], fill=confirmed_fill),
    )

    terms = workbook.create_sheet("术语表")
    terms.append(["概念", "中文建议", "英文建议", "适用范围", "备注"])
    term_rows = [
        ["安全格操作", "翻开", "Reveal", "棋盘直接操作", "不要与章节层面的“净化”混用"],
        ["章节目标", "净化", "Purify", "生态、关卡、章节完成", ""],
        ["危险对象", "污染核心", "Pollution Core", "通用、陆地、天空", "海洋机制内可称水雷"],
        ["标记动作", "标记", "Flag", "所有章节", "检测种子/珊瑚/纸飞机是视觉表现"],
        ["数字操作", "快速展开", "Quick Expand", "平面与宇宙棋盘", "鼠标为双击，键盘为 J"],
        ["棋盘单位", "格 / 面 / 星核", "Cell / Face / Star Core", "按章节固定", "同一界面不要混用"],
        ["重试同关", "重新开始", "Restart", "保持关卡编号", "是否重随机由玩法明确"],
        ["生成新局", "重新生成", "Regenerate", "生成新布局", "与 Restart 区分"],
    ]
    for row in term_rows:
        terms.append(row)
    for cell in terms[1]:
        cell.fill = header_fill
        cell.font = header_font
    for column, width in enumerate([18, 22, 24, 26, 48], 1):
        terms.column_dimensions[get_column_letter(column)].width = width
    for row in terms.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    terms.freeze_panes = "A2"

    guide = workbook.create_sheet("使用说明")
    guide.column_dimensions["A"].width = 110
    instructions = [
        "《绿色清扫者》文案维护说明",
        "1. “代号”是程序和翻译文件使用的永久键，不得修改、删除或重复。",
        "2. 人工精修时只修改“中文”和“English”两列，可填写修改备注并更新审核状态。",
        "3. 必须原样保留“占位符”列列出的 %d、%s、%02d、%.1f 等标记，数量和顺序都不能改变。",
        "4. “使用位置”只是辅助定位；同步依赖代号，不依赖源码行号，所以代码增删行不会造成文案错位。",
        "5. 修改完成后运行：python tools/copy_catalog.py import",
        "6. 导入命令会先验证代号和占位符，再同步 copy_catalog.json、中文 PO 和英文 PO。",
        "7. 如需根据仓库当前内容重新生成 Excel，运行：python tools/copy_catalog.py export。导出前先确保 Excel 修改已经 import。",
        "8. 提交或发布前运行：python tools/copy_catalog.py validate。",
        "9. 状态为“旧版、开发、非独立显示”的词条默认不属于当前玩家主流程，但仍保留永久代号，避免误删造成兼容问题。",
    ]
    for index, text in enumerate(instructions, 1):
        guide.cell(index, 1, text)
        guide.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="top")
    guide.cell(1, 1).font = Font(size=16, bold=True, color="31583F")

    metadata = workbook.create_sheet("元数据")
    metadata.sheet_state = "hidden"
    metadata.append(["schema", 1])
    metadata.append(["id_format", "TXT_####"])
    metadata.append(["entry_count", len(entries)])

    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(WORKBOOK)


def initialize() -> None:
    if MANIFEST.exists():
        raise SystemExit(f"Refusing to initialize: {MANIFEST} already exists")
    legacy_entries = parse_po(EN_PO)
    if not legacy_entries or any(ID_PATTERN.match(item["msgid"]) for item in legacy_entries):
        raise SystemExit("English PO no longer contains legacy Chinese msgids")

    additions = [
        {"msgid": "GREEN SWEEPER", "msgstr": "GREEN SWEEPER", "line": 0},
        {"msgid": "English", "msgstr": "English", "line": 0},
    ]
    source_entries = legacy_entries + additions
    values = {item["msgid"] for item in source_entries}
    occurrences = collect_occurrences(values)
    entries: list[dict[str, Any]] = []
    mapping: dict[str, str] = {}

    for index, source in enumerate(source_entries, 1):
        copy_id = f"TXT_{index:04d}"
        zh = source["msgid"]
        en = source["msgstr"]
        category = (
            "品牌与主菜单"
            if zh == "GREEN SWEEPER"
            else "设置"
            if zh == "English"
            else category_for_line(source["line"], zh)
        )
        entry = {
            "id": copy_id,
            "category": category,
            "status": status_for_text(zh),
            "zh_CN": zh,
            "en": en,
            "placeholders": placeholders(zh),
            "sources": occurrences.get(zh, []),
            "note": "",
            "review": "待复核",
        }
        if placeholders(zh) != placeholders(en):
            raise SystemExit(
                f"Placeholder mismatch before migration: {copy_id} {zh!r} / {en!r}"
            )
        entries.append(entry)
        mapping[zh] = copy_id

    replacement_counts = replace_source_literals(mapping)
    missing_usage = [
        entry["id"]
        for entry in entries
        if not replacement_counts.get(entry["id"]) and entry["status"] == "运行"
    ]
    if missing_usage:
        print(
            "Warning: runtime entries without a replaced source literal: "
            + ", ".join(missing_usage),
            file=sys.stderr,
        )

    refresh_source_locations(entries)
    save_manifest(entries)
    write_po(ZH_PO, "zh_CN", entries)
    write_po(EN_PO, "en", entries)
    export_workbook(entries)
    print(f"Initialized {len(entries)} stable copy IDs")
    print(f"Workbook: {WORKBOOK}")


def import_workbook() -> None:
    entries = load_manifest()
    by_id = {entry["id"]: entry for entry in entries}
    workbook = load_workbook(WORKBOOK)
    sheet = workbook["文案总表"]
    headers = {cell.value: index for index, cell in enumerate(sheet[1], 1)}
    required = {"代号", "分类", "状态", "中文", "English", "修改备注", "审核状态"}
    missing_headers = required - set(headers)
    if missing_headers:
        raise SystemExit(f"Workbook is missing columns: {sorted(missing_headers)}")

    seen: set[str] = set()
    for row in range(2, sheet.max_row + 1):
        copy_id = str(sheet.cell(row, headers["代号"]).value or "").strip()
        if not copy_id:
            continue
        if copy_id in seen:
            raise SystemExit(f"Duplicate copy ID in workbook: {copy_id}")
        if copy_id not in by_id:
            raise SystemExit(f"Unknown copy ID in workbook: {copy_id}")
        seen.add(copy_id)
        entry = by_id[copy_id]
        zh = str(sheet.cell(row, headers["中文"]).value or "")
        en = str(sheet.cell(row, headers["English"]).value or "")
        expected = entry["placeholders"]
        if placeholders(zh) != expected:
            raise SystemExit(
                f"Chinese placeholders changed for {copy_id}: "
                f"expected {expected}, got {placeholders(zh)}"
            )
        if placeholders(en) != expected:
            raise SystemExit(
                f"English placeholders changed for {copy_id}: "
                f"expected {expected}, got {placeholders(en)}"
            )
        entry["zh_CN"] = zh
        entry["en"] = en
        entry["category"] = str(sheet.cell(row, headers["分类"]).value or entry["category"])
        entry["status"] = str(sheet.cell(row, headers["状态"]).value or entry["status"])
        entry["note"] = str(sheet.cell(row, headers["修改备注"]).value or "")
        entry["review"] = str(sheet.cell(row, headers["审核状态"]).value or "待复核")

    missing_ids = set(by_id) - seen
    if missing_ids:
        raise SystemExit(f"Workbook is missing copy IDs: {sorted(missing_ids)}")

    save_manifest(entries)
    write_po(ZH_PO, "zh_CN", entries)
    write_po(EN_PO, "en", entries)
    print(f"Imported {len(entries)} rows and updated both PO catalogs")


def validate() -> None:
    entries = load_manifest()
    errors: list[str] = []
    manifest_ids = [entry["id"] for entry in entries]
    if len(manifest_ids) != len(set(manifest_ids)):
        errors.append("Manifest contains duplicate IDs")
    for entry in entries:
        if not ID_PATTERN.match(entry["id"]):
            errors.append(f"Invalid ID format: {entry['id']}")
        expected = entry["placeholders"]
        if placeholders(entry["zh_CN"]) != expected:
            errors.append(f"Chinese placeholder mismatch: {entry['id']}")
        if placeholders(entry["en"]) != expected:
            errors.append(f"English placeholder mismatch: {entry['id']}")

    for path, locale in [(ZH_PO, "zh_CN"), (EN_PO, "en")]:
        po_entries = parse_po(path)
        po_ids = [entry["msgid"] for entry in po_entries]
        if po_ids != manifest_ids:
            errors.append(f"{path.name} ID order/set differs from manifest")
        expected_text = {entry["id"]: entry[locale] for entry in entries}
        for po_entry in po_entries:
            if expected_text.get(po_entry["msgid"]) != po_entry["msgstr"]:
                errors.append(f"{path.name} text differs for {po_entry['msgid']}")

    source_ids: set[str] = set()
    raw_chinese_literals: list[str] = []
    for path in iter_source_files():
        text = path.read_text(encoding="utf-8")
        relative = path.relative_to(ROOT).as_posix()
        for match in QUOTED_STRING_PATTERN.finditer(text):
            value = decode_quoted(match.group(0))
            if not isinstance(value, str):
                continue
            if ID_PATTERN.match(value):
                source_ids.add(value)
            elif HAN_PATTERN.search(value):
                line = text.count("\n", 0, match.start()) + 1
                raw_chinese_literals.append(f"{relative}:{line} {value!r}")
    unknown_source_ids = source_ids - set(manifest_ids)
    if unknown_source_ids:
        errors.append(f"Unknown IDs in source: {sorted(unknown_source_ids)}")
    if raw_chinese_literals:
        errors.append(
            "Raw Chinese literals bypass stable copy IDs: "
            + "; ".join(raw_chinese_literals)
        )

    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
    print(
        f"Validated {len(entries)} IDs, both PO catalogs, placeholders, and source references"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["initialize", "export", "import", "validate"])
    args = parser.parse_args()
    if args.command == "initialize":
        initialize()
    elif args.command == "export":
        entries = load_manifest()
        refresh_source_locations(entries)
        save_manifest(entries)
        export_workbook(entries)
        print(f"Exported {len(entries)} rows to {WORKBOOK}")
    elif args.command == "import":
        import_workbook()
    else:
        validate()


if __name__ == "__main__":
    main()
